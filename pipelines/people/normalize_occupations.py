import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
JSON_FILE = ROOT / "data" / "people" / "occupations.json"
MAPPING_FILE = ROOT / "data" / "people" / "mappings" / "to_replace.xlsx"


def replace_occup_by_2_col(json_file_path: Path, excel_file_path: Path) -> None:
    with json_file_path.open("r", encoding="utf-8") as json_file:
        data = json.load(json_file)

    workbook = openpyxl.load_workbook(excel_file_path)
    sheet = workbook.active
    replace_dict = {
        sheet.cell(row=row, column=1).value: sheet.cell(row=row, column=2).value
        for row in range(1, sheet.max_row + 1)
    }

    for item in data:
        if item.get("occupations") in replace_dict:
            item["occupations"] = replace_dict[item["occupations"]]

    with json_file_path.open("w", encoding="utf-8") as json_file:
        json.dump(data, json_file, indent=2, ensure_ascii=False)


if __name__ == "__main__":
    replace_occup_by_2_col(JSON_FILE, MAPPING_FILE)
